import io
import pytest
import docx
from httpx import AsyncClient, ASGITransport
from app.main import app
from app.db.database import AsyncSessionLocal
from app.db.init_db import init_tables
from app.services.document_processor import DocumentProcessor
from app.services.vector_store import vector_store
from app.ai_engine.memory import TieredMemory
from app.ai_engine.providers.local_provider import LocalLLMProvider

@pytest.mark.asyncio
async def test_text_splitter_logic():
    """Valida que el particionador de texto respete límites y solapamientos."""
    sample_text = (
        "Primer párrafo sobre la arquitectura local de Lucil AI. "
        "Segundo párrafo con especificaciones de hardware y memoria GPU. "
        "Tercer párrafo detallando el protocolo de cero alucinación."
    )
    chunks = DocumentProcessor.split_text(sample_text, chunk_size=80, chunk_overlap=20)
    assert len(chunks) >= 2
    for chunk in chunks:
        assert len(chunk) <= 90  # Con margen de límites naturales

@pytest.mark.asyncio
async def test_vector_store_and_tiered_memory_rag():
    """Valida la indexación y recuperación semántica en ChromaDB con nomic-embed-text."""
    test_user_id = "test_rag_user_uuid_123"
    test_project_id = "proj_456"
    test_filename = "manual_hardware_amd.docx"

    # 1. Ingestar fragmentos de prueba
    doc_text = (
        "La tarjeta gráfica AMD Radeon RX 9060 XT cuenta con 16 GB de VRAM dedicada. "
        "Utiliza la arquitectura RDNA 4 / ROCm v7.1 gfx1200 para inferencia ultra rápida. "
        "Permite ejecutar Qwen 2.5 14B con 32k tokens de contexto totalmente en memoria local."
    )
    success = await DocumentProcessor.chunk_and_embed(
        text=doc_text,
        project_id=test_project_id,
        filename=test_filename,
        user_id=test_user_id
    )
    assert success is True

    # 2. Búsqueda semántica directa en VectorStoreService
    results = await vector_store.similarity_search(
        query="Cuánta memoria VRAM tiene la GPU de AMD?",
        n_results=2,
        where={"user_id": test_user_id}
    )
    assert len(results) > 0
    top_doc = results[0]
    assert "16 GB de VRAM" in top_doc["content"]
    assert top_doc["metadata"]["filename"] == test_filename

    # 3. Recuperación a través de TieredMemory
    tiered_memory = TieredMemory()
    long_term_results = await tiered_memory.retrieve_long_term_memory(
        query="arquitectura y tokens de contexto",
        user_id=test_user_id,
        project_id=test_project_id
    )
    assert len(long_term_results) > 0
    assert any("contexto" in r["content"] for r in long_term_results)

    # 4. Inyección en el prompt de LocalLLMProvider
    local_provider = LocalLLMProvider()
    context = {
        "semantic": {"preferencias": ["Usar siempre español"]},
        "long_term": long_term_results,
        "short_term": []
    }
    messages = local_provider._build_messages(
        prompt="Qué características tiene mi tarjeta?",
        context=context
    )
    system_msg = messages[0]["content"]
    assert "[Información Recuperada de tus Documentos en Memoria]:" in system_msg
    assert f"[Documento propio: {test_filename}]" in system_msg
    assert "16 GB de VRAM" in system_msg

    # 5. Limpieza de vectores de prueba
    await vector_store.delete_by_metadata(where={"user_id": test_user_id})

@pytest.mark.asyncio
async def test_document_upload_endpoint():
    """Valida la subida real de un documento docx a través de la API REST."""
    await init_tables()
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        # Autenticación
        email = "test_rag_upload@lucil.ai"
        password = "Password123!"
        await client.post("/api/v1/auth/register", json={"email": email, "password": password})
        login_resp = await client.post(
            "/api/v1/auth/login",
            data={"username": email, "password": password},
            headers={"Content-Type": "application/x-www-form-urlencoded"}
        )
        assert login_resp.status_code == 200
        token = login_resp.json()["access_token"]
        auth_headers = {"Authorization": f"Bearer {token}"}

        # Generar un archivo docx en memoria
        doc = docx.Document()
        doc.add_paragraph("Tratamiento farmacológico de referencia: Protocolo de prueba para cefalea tensional.")
        doc.add_paragraph("El principio activo sugerido por la guía es ibuprofeno 400 mg con las comidas.")
        bio = io.BytesIO()
        doc.save(bio)
        bio.seek(0)

        # Subir documento
        files = {
            "file": ("guia_clinica.docx", bio.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        }
        resp = await client.post(
            "/api/v1/documents/proj_salud_01/upload",
            files=files,
            headers=auth_headers
        )
        assert resp.status_code == 200
        data = resp.json()
        assert "guia_clinica.docx indexado exitosamente" in data["message"]